"""
tests/test_markdown_convert.py - markitdown 檔案轉 Markdown 前處理器測試

涵蓋：docx/xlsx/pptx 轉換、內嵌圖片濾除、僅限本機檔案、不啟用外掛與外部 LLM、錯誤包裝。
markitdown 為選用依賴，未安裝時整個檔案略過。
"""

import struct
import zlib

import pytest

pytest.importorskip("markitdown")

import docx  # noqa: E402
import openpyxl  # noqa: E402

from src.ingestion.markdown_convert import (  # noqa: E402
    MarkdownConvertError,
    convert_to_markdown,
    strip_embedded_images,
)


def _tiny_png() -> bytes:
    def chunk(tag: bytes, data: bytes) -> bytes:
        c = struct.pack(">I", len(data)) + tag + data
        return c + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    raw = b"\x00\xff\x00\x00"
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def test_docx_paragraph_and_table(tmp_path):
    d = docx.Document()
    d.add_paragraph("甲溝炎是常見的足部問題")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "項目", "說明"
    t.cell(1, 0).text, t.cell(1, 1).text = "冰敷", "每次十分鐘"
    p = tmp_path / "a.docx"
    d.save(p)

    md = convert_to_markdown(p)
    assert "甲溝炎是常見的足部問題" in md
    assert "| 項目" in md and "冰敷" in md


def test_docx_embedded_image_removed(tmp_path):
    img = tmp_path / "x.png"
    img.write_bytes(_tiny_png())
    d = docx.Document()
    d.add_paragraph("含圖片的文件")
    d.add_picture(str(img))
    p = tmp_path / "img.docx"
    d.save(p)

    md = convert_to_markdown(p)
    assert "含圖片的文件" in md
    assert "data:image" not in md and "base64" not in md and "![" not in md


def test_xlsx_to_markdown_table(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "客服"
    ws.append(["問題", "回覆"])
    ws.append(["營業時間", "請致電診所確認"])
    p = tmp_path / "a.xlsx"
    wb.save(p)

    md = convert_to_markdown(p)
    assert "營業時間" in md and "|" in md


def test_pptx_converts(tmp_path):
    pptx = pytest.importorskip("pptx")
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "脂肪瘤手術方式"
    slide.placeholders[1].text = "傳統切除術與微創切除術"
    p = tmp_path / "a.pptx"
    prs.save(p)

    md = convert_to_markdown(p)
    assert "脂肪瘤手術方式" in md and "微創切除術" in md


@pytest.mark.parametrize("name", ["a.txt", "a.jpg", "a.doc"])
def test_unsupported_suffix_rejected(tmp_path, name):
    p = tmp_path / name
    p.write_bytes(b"x")
    with pytest.raises(MarkdownConvertError, match="不支援"):
        convert_to_markdown(p)


def test_missing_file_rejected(tmp_path):
    with pytest.raises(MarkdownConvertError, match="找不到"):
        convert_to_markdown(tmp_path / "nope.docx")


@pytest.mark.parametrize("src", ["http://example.com/a.docx", "https://x.y/a.pdf", "data:text/plain;base64,AAAA", "file:///etc/a.docx"])
def test_only_local_paths_allowed(src):
    with pytest.raises(MarkdownConvertError, match="本機"):
        convert_to_markdown(src)


def test_empty_document_rejected(tmp_path):
    p = tmp_path / "empty.docx"
    docx.Document().save(p)
    with pytest.raises(MarkdownConvertError, match="文字層"):
        convert_to_markdown(p)


def test_corrupt_file_wrapped(tmp_path):
    p = tmp_path / "~$lock.docx"
    p.write_bytes(b"x" * 162)
    with pytest.raises(MarkdownConvertError):
        convert_to_markdown(p)


def test_no_plugins_no_llm(monkeypatch, tmp_path):
    """純本地推理：不得啟用外掛，也不得傳入任何 LLM／雲端客戶端。"""
    import markitdown

    seen = {}

    class Recorder:
        def __init__(self, *args, **kwargs):
            seen["args"], seen["kwargs"] = args, kwargs

        def convert(self, path):
            class R:
                text_content = "內容"

            return R()

    monkeypatch.setattr(markitdown, "MarkItDown", Recorder)
    p = tmp_path / "a.docx"
    docx.Document().save(p)
    assert convert_to_markdown(p) == "內容"
    assert seen["kwargs"].get("enable_plugins") is False
    for banned in ("llm_client", "llm_model", "docintel_endpoint"):
        assert banned not in seen["kwargs"]


def test_strip_embedded_images_unit():
    md = "前\n\n![圖](data:image/jpeg;base64,AAAA)\n\n![C:\\x.jpg](data:image/png;base64...)\n\n\n\n後"
    out = strip_embedded_images(md)
    assert "data:" not in out and "![" not in out
    assert out.startswith("前") and out.endswith("後")
    assert "\n\n\n" not in out
