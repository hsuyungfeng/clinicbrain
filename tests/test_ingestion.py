#!/usr/bin/env python3
"""
文件文字擷取、簡繁轉換與 FAQ 生成驗證層單元測試
Phase 03 Document Ingestion Stage 1: TASK-01, TASK-02, TASK-03
"""

import json
from pathlib import Path
import pytest
import docx
import openpyxl

from src.ingestion.convert_chinese import to_traditional
from src.ingestion.extract_text import extract_docx_text, extract_xlsx_text
from src.ingestion.generate_faq import (
    build_faq_prompt,
    parse_and_validate_faq,
    validate_single_faq,
)

PROJECT_ROOT = Path(__file__).resolve().parent.parent
ORIGINAL_DATA_DIR = PROJECT_ROOT / "OriginalData" / "緻妍外科診所"


# =========================================================================
# TASK-01: 文字擷取測試
# =========================================================================

def test_extract_docx_paragraphs_and_tables(tmp_path: Path):
    """驗證 extract_docx_text 同時擷取段落與表格文字（不漏掉表格資訊）。"""
    test_docx_path = tmp_path / "test_doc.docx"
    doc = docx.Document()
    doc.add_paragraph("這是第一段門診衛教文字。")
    doc.add_paragraph("這是第二段手術說明。")
    
    # 加入表格
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "項目"
    table.cell(0, 1).text = "費用說明"
    table.cell(1, 0).text = "甲溝炎手術"
    table.cell(1, 1).text = "自費處理請洽診所"
    
    doc.save(str(test_docx_path))

    extracted = extract_docx_text(test_docx_path)
    assert "這是第一段門診衛教文字。" in extracted
    assert "這是第二段手術說明。" in extracted
    assert "甲溝炎手術" in extracted
    assert "自費處理請洽診所" in extracted
    assert "--- 表格 1 ---" in extracted


def test_extract_xlsx_sheets_and_rows(tmp_path: Path):
    """驗證 extract_xlsx_text 忠實擷取各工作表欄位為字典結構。"""
    test_xlsx_path = tmp_path / "test_sheet.xlsx"
    wb = openpyxl.Workbook()
    
    # 工作表 1
    ws1 = wb.active
    ws1.title = "客戶諮詢"
    ws1.append(["問題", "回覆話術"])
    ws1.append(["有沒有體驗價？", "目前開幕期間有專案優惠"])
    
    # 工作表 2
    ws2 = wb.create_sheet(title="術後關懷")
    ws2.append(["階段", "話術內容"])
    ws2.append(["術後第1天", "請問回家後傷口照護還順利嗎？"])
    
    wb.save(str(test_xlsx_path))

    extracted = extract_xlsx_text(test_xlsx_path)
    assert "客戶諮詢" in extracted
    assert "術後關懷" in extracted
    assert len(extracted["客戶諮詢"]) == 1
    assert extracted["客戶諮詢"][0]["問題"] == "有沒有體驗價？"
    assert extracted["術後關懷"][0]["話術內容"] == "請問回家後傷口照護還順利嗎？"


@pytest.mark.skipif(not ORIGINAL_DATA_DIR.exists(), reason="找不到 OriginalData 目錄")
def test_extract_priority_files_completeness():
    """驗證三份真實優先檔案的擷取結果完整性與非空性。"""
    docx_file1 = ORIGINAL_DATA_DIR / "緻妍自費門診手術內容.docx"
    docx_file2 = ORIGINAL_DATA_DIR / "緻妍可自費門診手術 可配合開立診斷書.docx"
    xlsx_file = ORIGINAL_DATA_DIR / "客服回覆話術.xlsx"

    text1 = extract_docx_text(docx_file1)
    assert len(text1) > 1000, f"docx1 擷取字數過少: {len(text1)}"
    # 確保 5 張表格有被讀取到
    assert "--- 表格 1 ---" in text1
    assert "--- 表格 5 ---" in text1

    text2 = extract_docx_text(docx_file2)
    assert len(text2) > 1000, f"docx2 擷取字數過少: {len(text2)}"

    sheets = extract_xlsx_text(xlsx_file)
    assert len(sheets) == 4
    assert "工作表1" in sheets
    assert "客戶詢問" in sheets
    assert len(sheets["客戶詢問"]) > 10


# =========================================================================
# TASK-02: 簡繁轉換測試 (opencc s2twp)
# =========================================================================

def test_convert_chinese_s2twp_idioms():
    """驗證 s2twp 簡繁轉換正確性與台灣慣用詞本地化。"""
    # 簡體到台灣繁體
    assert to_traditional("这是一个医疗软件") == "這是一個醫療軟體"
    assert to_traditional("手术后屏幕显示信息") == "手術後螢幕顯示資訊"
    assert to_traditional("电脑内存") == "電腦記憶體"


def test_convert_chinese_idempotent():
    """驗證已是繁體中文的文字在轉換後保持不變（冪等性）。"""
    traditional_text = "緻妍外科診所提供自費門診手術與術後照護衛教。"
    assert to_traditional(traditional_text) == traditional_text
    assert to_traditional("") == ""


# =========================================================================
# TASK-03: FAQ 驗證與過濾層測試
# =========================================================================

def test_validate_real_price_leakage_rejection():
    """驗證來源文件真實出現過的多種價格格式均會被嚴格攔截剔除。"""
    real_prices = [
        "$1500",
        "$4500",
        "$6000",
        "$8000",
        "$200",
        "$2000",
        "1500元",
        "NT$ 500",
        "NT$1200",
        "健保點值X2.5",
        "以健保點值*2.5計算",
    ]

    for price in real_prices:
        faq = {
            "question": "甲溝炎手術收費如何？",
            "answer": f"門診處理費用大約為 {price}，請於門診時諮詢。",
        }
        is_valid, reason = validate_single_faq(faq)
        assert not is_valid, f"價格 '{price}' 未被成功攔截！"
        assert "具體金額或價格數字" in reason


def test_validate_political_stance_rejection():
    """驗證政治立場與主權表述詞彙被攔截。"""
    forbidden_political = [
        "台灣是不可分割的一部分",
        "在一個中國原則下進行診療",
        "中國台灣地區的診所規範",
        "本院遵守台灣地區醫療法令",
    ]
    for phrase in forbidden_political:
        faq = {
            "question": "診所服務範圍？",
            "answer": f"服務說明：{phrase}。",
        }
        is_valid, reason = validate_single_faq(faq)
        assert not is_valid, f"政治立場詞彙 '{phrase}' 未被攔截！"
        assert "政治立場詞彙" in reason


def test_validate_simplified_chinese_rejection():
    """驗證簡體中文字元被攔截。"""
    faq = {
        "question": "术后如何护理？",
        "answer": "这是关于门诊手术的护理说明。",
    }
    is_valid, reason = validate_single_faq(faq)
    assert not is_valid, "簡體中文字元未被攔截！"
    assert "簡體中文字符" in reason


def test_validate_forbidden_guarantee_rejection():
    """驗證保證療效語句被攔截。"""
    faq = {
        "question": "疤痕修復效果好嗎？",
        "answer": "醫師施打消疤針保證有效消除疤痕組織。",
    }
    is_valid, reason = validate_single_faq(faq)
    assert not is_valid, "保證療效詞彙未被攔截！"
    assert "違規禁詞" in reason


def test_parse_and_validate_faq_mixed_batch():
    """驗證批次驗證時，違規項目被剔除但不影響合格項目的保留。"""
    mock_llm_json = json.dumps([
        {
            "question": "粉瘤切除會留疤嗎？",
            "answer": "手術切口會依皮膚紋理設計以盡量減少痕跡，費用請致電診所確認。",
        },
        {
            "question": "甲溝炎收費多少？",
            "answer": "自費處理傷口一處收費$4500元，雙側同趾收費$6000元。",  # 含真實價格，應被剔除
        },
        {
            "question": "術後可以開立診斷證明書嗎？",
            "answer": "可於門診時告知醫師，由醫師評估後開立門診手術診斷證明書以利保險申請。",
        },
        {
            "question": "手术需要拆线吗？",  # 簡體字，應被剔除
            "answer": "一般缝合于7至10天后回诊由医师评估拆线。",
        },
    ], ensure_ascii=False)

    valid_faqs, rejected_faqs = parse_and_validate_faq(mock_llm_json, return_rejected=True)
    assert len(valid_faqs) == 2, f"預期保留 2 筆合格 FAQ，實際保留 {len(valid_faqs)}"
    assert len(rejected_faqs) == 2, f"預期剔除 2 筆違規 FAQ，實際剔除 {len(rejected_faqs)}"

    assert valid_faqs[0]["question"] == "粉瘤切除會留疤嗎？"
    assert valid_faqs[1]["question"] == "術後可以開立診斷證明書嗎？"

    rejected_reasons = [r["reason"] for r in rejected_faqs]
    assert any("具體金額" in r for r in rejected_reasons)
    assert any("簡體中文" in r for r in rejected_reasons)


def test_build_faq_prompt_contains_constraints():
    """驗證 build_faq_prompt 輸出的提示詞明確包含價格屏蔽與繁體中文約束。"""
    prompt = build_faq_prompt("測試內容文字", "測試檔案.docx")
    assert "繁體中文" in prompt
    assert "價格屏蔽" in prompt
    assert "請致電診所確認" in prompt
    assert "測試檔案.docx" in prompt
