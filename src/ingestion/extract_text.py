#!/usr/bin/env python3
"""
Taiwan Clinic Medical PageIndex RAG System - 文件文字擷取模組 (extract_text)
Phase 03 Document Ingestion Stage 1: TASK-01

支援格式：
- docx: 同時擷取段落 (paragraphs) 與表格 (tables)，確保手術報價、收費項目不遺漏。
- xlsx: 忠實擷取工作表每一列資料為結構化字典，不預設欄位語意，交由 LLM 判斷。
"""

from pathlib import Path
from typing import Any, Union
import docx
import openpyxl


def extract_docx_text(path: Union[str, Path]) -> str:
    """讀取 .docx 檔案，完整擷取段落與表格文字。

    參數:
        path: .docx 檔案之路徑
    回傳:
        結合段落與表格內容之純文字字串。
    """
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"找不到 docx 檔案: {file_path}")

    doc = docx.Document(str(file_path))
    content_parts: list[str] = []

    # 1. 擷取所有段落文字
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            content_parts.append(text)

    # 2. 擷取所有表格文字（避免手術價格或說明漏掉）
    for table_idx, table in enumerate(doc.tables):
        table_lines: list[str] = [f"--- 表格 {table_idx + 1} ---"]
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            # 若整列皆為空則略過
            if any(row_cells):
                table_lines.append("\t".join(row_cells))
        if len(table_lines) > 1:
            content_parts.append("\n".join(table_lines))

    return "\n\n".join(content_parts)


def extract_xlsx_text(path: Union[str, Path]) -> dict[str, list[dict[str, Any]]]:
    """讀取 .xlsx 檔案，回傳各工作表的列資料字典清單。

    參數:
        path: .xlsx 檔案之路徑
    回傳:
        字典 {工作表名稱: [列資料字典, ...]}
    """
    file_path = Path(path)
    if not file_path.exists():
        raise FileNotFoundError(f"找不到 xlsx 檔案: {file_path}")

    wb = openpyxl.load_workbook(str(file_path), data_only=True)
    result: dict[str, list[dict[str, Any]]] = {}

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            result[sheet_name] = []
            continue

        # 第一列視為欄位標題
        raw_headers = rows[0]
        headers: list[str] = []
        for idx, h in enumerate(raw_headers):
            if h is not None and str(h).strip():
                headers.append(str(h).strip())
            else:
                headers.append(f"col_{idx + 1}")

        sheet_records: list[dict[str, Any]] = []
        for row in rows[1:]:
            # 檢查整列是否皆為空
            if not any(v is not None and str(v).strip() for v in row):
                continue

            record: dict[str, Any] = {}
            for col_idx, header in enumerate(headers):
                val = row[col_idx] if col_idx < len(row) else None
                if val is not None:
                    str_val = str(val).strip()
                    if str_val:
                        record[header] = str_val
            if record:
                sheet_records.append(record)

        result[sheet_name] = sheet_records

    return result
